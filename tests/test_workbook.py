import copy
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills' / 'frcstack' / 'scripts'
sys.path.insert(0, str(SCRIPTS))

def event():
    return {'schema': 1, 'program': 'FRC', 'season': '2026', 'event_id': 'demo',
            'teams': [{'number': '111', 'name': 'Synthetic Alpha'}, {'number': '222', 'name': 'Synthetic Beta'}],
            'sources': [{'url': 'https://example.org/demo', 'description': 'Synthetic test, not real scouting'}]}

def observation(rid, kind, start, end, **extra):
    return dict(id=rid, team='111', kind=kind, start_s=start, end_s=end,
                visibility='direct', reviewed=True, is_replay=False, evidence='Synthetic frame evidence', **extra)

def ledger():
    return {'schema': 1, 'program': 'FRC', 'season': '2026', 'event_id': 'demo', 'match_id': 'q1',
            'source_url': 'https://example.org/demo', 'observations': [
                observation('a', 'intake_window', 0, 2, count=10),
                observation('b', 'intake_window', 2, 6, count=2),
                observation('s', 'score', 7, 7, outcome='scored', points=2, rule_ref='Synthetic rule')
            ]}

class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = SCRIPTS / 'export_scouting.py'
        assert path.is_file(), 'Workbook exporter is not implemented'
        spec = importlib.util.spec_from_file_location('export_scouting', path)
        cls.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.module)

    def export(self, e=None, ledgers=None):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / 'scouting.xlsx'
        self.module.export_event(e or event(), ledgers if ledgers is not None else [ledger()], path)
        book = load_workbook(path)
        self.addCleanup(book.close)
        return path, book

    def rows(self, sheet):
        values = list(sheet.values)
        return [dict(zip(values[0], row)) for row in values[1:]]

    def test_all_roster_teams_and_weighted_rate(self):
        _, book = self.export()
        self.assertEqual(book.sheetnames, ['Teams', 'Matches', 'Observations', 'Mechanisms', 'Recommendations', 'Sources', 'Readme'])
        rows = self.rows(book['Teams'])
        self.assertEqual([r['Team'] for r in rows], ['111', '222'])
        self.assertEqual(rows[0]['Intake pieces/s'], 2)
        self.assertEqual(rows[0]['Observed points (lower bound)'], 2)
        self.assertEqual(rows[1]['Intake pieces/s'], 'Unknown')
        self.assertEqual(rows[1]['Observed scoring events'], 'Unknown')
        self.assertEqual(book['Teams'].freeze_panes, 'C2')
        self.assertTrue(book['Teams'].tables)

    def test_empty_capture_leaves_unknown_and_team_row(self):
        _, book = self.export(ledgers=[])
        self.assertEqual(len(self.rows(book['Teams'])), 2)
        self.assertEqual(self.rows(book['Teams'])[0]['Observed matches'], 0)
        self.assertEqual(self.rows(book['Teams'])[0]['Observed scoring events'], 'Unknown')

    def test_replays_corrections_occlusion_and_missing_points(self):
        data = ledger()
        data['observations'] = [
            observation('a', 'intake_window', 0, 2, count=99),
            observation('b', 'intake_window', 0, 2, count=4, supersedes='a'),
            observation('r', 'fire_window', 3, 4, count=100),
            observation('x', 'score', 5, 5, outcome='scored')]
        data['observations'][2]['is_replay'] = True
        data['observations'][3]['visibility'] = 'occluded'
        _, book = self.export(ledgers=[data])
        row = self.rows(book['Teams'])[0]
        self.assertEqual(row['Intake pieces/s'], 2)
        self.assertEqual(row['Firing pieces/s'], 'Unknown')
        self.assertEqual(row['Observed points (lower bound)'], 'Unknown')
        self.assertEqual(book['Observations'].max_row, 5)

    def test_rejects_inconsistent_event_duplicate_match_and_nonroster_team(self):
        for field, value in [('season', '2025'), ('event_id', 'other'), ('program', 'FTC')]:
            data = ledger(); data[field] = value
            with self.assertRaises(ValueError): self.export(ledgers=[data])
        with self.assertRaises(ValueError): self.export(ledgers=[ledger(), ledger()])
        data = ledger(); data['observations'][0]['team'] = '999'
        with self.assertRaises(ValueError): self.export(ledgers=[data])

    def test_text_is_literal_not_a_formula(self):
        data = event(); data['teams'][0]['name'] = '=HYPERLINK("https://bad.invalid","click")'
        source = ledger(); source['observations'][0]['evidence'] = '+malicious text'
        _, book = self.export(e=data, ledgers=[source])
        self.assertEqual(book['Teams']['B2'].value, data['teams'][0]['name'])
        self.assertEqual(book['Teams']['B2'].data_type, 's')
        self.assertFalse(any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row))

    def test_existing_output_and_malformed_roster_refused(self):
        path, _ = self.export()
        with self.assertRaises(FileExistsError): self.module.export_event(event(), [ledger()], path)
        data = event(); data['teams'].append(copy.deepcopy(data['teams'][0]))
        with self.assertRaises(ValueError): self.export(e=data)

if __name__ == '__main__': unittest.main()
