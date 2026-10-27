#!/usr/bin/env python3
"""Export reviewed observation ledgers to an evidence-preserving XLSX snapshot; no video ingestion."""
import argparse
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from summarize_scouting import summarize

UNKNOWN = 'Unknown'

def text(value):
    if value is None:
        return UNKNOWN
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    return value

def add_sheet(book, title, headers, rows):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([text(v) for v in row])
    for row in sheet:
        for cell in row:
            # External strings stay strings, even if they begin with '='. Never evaluate captured text.
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError('Cell text exceeds Excel limit')
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=10, color='182B3A')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = '0.00'
    for cell in sheet[1]:
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='183C4A')
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = 'C2' if title == 'Teams' else 'B2'
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(55, max(18, len(header) + 2))
    if rows:
        table = Table(displayName='FRC' + title, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    sheet.auto_filter.ref = sheet.dimensions
    return sheet

def export_event(event: dict, ledgers: list[dict], output: Path) -> Path:
    output = Path(output)
    if output.exists():
        raise FileExistsError(str(output))
    if event.get('schema') != 1 or event.get('program') not in {'FRC', 'FTC'}:
        raise ValueError('Set actual schema, program, season and event identity')
    for key in ('season', 'event_id'):
        if not isinstance(event.get(key), str) or not event[key].strip():
            raise ValueError('Missing event field: ' + key)
    roster = event.get('teams')
    if not isinstance(roster, list) or not roster:
        raise ValueError('Provide the verified event roster')
    teams = {}
    for entry in roster:
        number = entry.get('number') if isinstance(entry, dict) else None
        if not isinstance(number, str) or not number.isascii() or not number.isdigit() or int(number) <= 0 or str(int(number)) != number or number in teams:
            raise ValueError('Team numbers must be unique canonical positive digit strings')
        name = entry.get('name')
        if name is not None and not isinstance(name, str):
            raise ValueError('Team name must be text')
        teams[number] = {'name': name, 'matches': 0, 'scores': None, 'points': None,
                         'intake_count': 0, 'intake_s': 0, 'fire_count': 0, 'fire_s': 0, 'unresolved': 0, 'replays': 0}
    match_ids = set()
    match_rows, observation_rows = [], []
    sources = list(event.get('sources', []))
    for ledger in ledgers:
        if any(ledger.get(k) != event[k] for k in ('program', 'season', 'event_id')):
            raise ValueError('Ledger identity does not match event')
        result = summarize(ledger)
        if result['match_id'] in match_ids:
            raise ValueError('Duplicate match ledger: merge/correct observations before export')
        match_ids.add(result['match_id'])
        if not set(result['teams']).issubset(teams):
            raise ValueError('Observed team is not in event roster')
        for number, summary in result['teams'].items():
            aggregate = teams[number]
            aggregate['matches'] += 1
            for destination, key in [('scores', 'reviewed_visible_scoring_events'), ('points', 'video_observed_points_lower_bound')]:
                if summary[key] is not None:
                    aggregate[destination] = (aggregate[destination] or 0) + summary[key]
            for prefix, key in [('intake', 'intake'), ('fire', 'firing')]:
                if summary[key] is not None:
                    aggregate[prefix + '_count'] += summary[key]['visible_count']
                    aggregate[prefix + '_s'] += summary[key]['observed_active_s']
            aggregate['unresolved'] += summary['unresolved_records']
            aggregate['replays'] += summary['excluded_replays']
            match_rows.append([ledger['match_id'], number, summary['reviewed_visible_scoring_events'], summary['video_observed_points_lower_bound'],
                               summary['intake']['pieces_per_s'] if summary['intake'] else None,
                               summary['firing']['pieces_per_s'] if summary['firing'] else None,
                               summary['unresolved_records'], summary['excluded_replays'], ledger['source_url']])
        superseded = {r['supersedes'] for r in ledger['observations'] if r.get('supersedes')}
        for record in ledger['observations']:
            status = ('Superseded' if record['id'] in superseded else 'Replay' if record['is_replay'] else
                      'Eligible' if record['visibility'] == 'direct' and record['reviewed'] else 'Unresolved')
            observation_rows.append([ledger['match_id'], record['team'], record['id'], record['kind'], record['start_s'], record['end_s'],
                                     record.get('count'), record.get('outcome'), record.get('points'), record.get('rule_ref'),
                                     record['visibility'], record['reviewed'], record['is_replay'], record.get('supersedes'), status,
                                     record['evidence'], ledger['source_url'], record])
        sources.append({'url': ledger['source_url'], 'description': 'Observation source for match ' + ledger['match_id']})
    team_rows = []
    for number, aggregate in sorted(teams.items(), key=lambda pair: int(pair[0])):
        team_rows.append([number, aggregate['name'], aggregate['matches'], aggregate['scores'], aggregate['points'],
                          aggregate['intake_count'] if aggregate['intake_s'] else None,
                          aggregate['intake_s'] or None,
                          aggregate['intake_count'] / aggregate['intake_s'] if aggregate['intake_s'] else None,
                          aggregate['fire_count'] if aggregate['fire_s'] else None,
                          aggregate['fire_s'] or None,
                          aggregate['fire_count'] / aggregate['fire_s'] if aggregate['fire_s'] else None,
                          aggregate['unresolved'], aggregate['replays'],
                          'Reviewed video lower bounds; no observations during camera gaps; see raw records and Sources'])
    book = Workbook(); book.remove(book.active)
    book.properties.title = 'FRCstack — ' + event['event_id']
    book.properties.description = 'Evidence snapshot, not official per-team scores or continuous telemetry'
    add_sheet(book, 'Teams', ['Team', 'Name', 'Observed matches', 'Observed scoring events', 'Observed points (lower bound)', 'Intake visible pieces', 'Intake active seconds', 'Intake pieces/s', 'Fired visible pieces', 'Firing active seconds', 'Firing pieces/s', 'Unresolved records', 'Excluded replays', 'Coverage / limitations'], team_rows)
    add_sheet(book, 'Matches', ['Match', 'Team', 'Observed scoring events', 'Observed points (lower bound)', 'Intake pieces/s', 'Firing pieces/s', 'Unresolved records', 'Excluded replays', 'Source'], match_rows)
    add_sheet(book, 'Observations', ['Match', 'Team', 'Record', 'Kind', 'Start seconds', 'End seconds', 'Count', 'Outcome', 'Points', 'Rule reference', 'Visibility', 'Reviewed', 'Replay', 'Supersedes', 'Status', 'Evidence', 'Source', 'Raw record'], observation_rows)
    for sheet, key, fields in [('Mechanisms', 'mechanisms', ['team', 'description', 'evidence', 'source_url', 'uncertainty']),
                               ('Recommendations', 'recommendations', ['team', 'action', 'evidence', 'source_url', 'bench_test', 'uncertainty'])]:
        rows = []
        for record in event.get(key, []):
            if record.get('team') not in teams or not record.get('evidence') or not record.get('source_url'):
                raise ValueError('Mechanisms/recommendations need roster team and source evidence')
            rows.append([record.get(field) for field in fields])
        add_sheet(book, sheet, [field.replace('_', ' ').title() for field in fields], rows)
    add_sheet(book, 'Sources', ['URL', 'Description'], [[s.get('url'), s.get('description')] for s in sources])
    add_sheet(book, 'Readme', ['Field', 'Meaning'], [
        ['Program / season / event', event['program'] + ' / ' + event['season'] + ' / ' + event['event_id']],
        ['Data status', event.get('data_status', 'Partial reviewed observation snapshot; no completeness claim')],
        ['Unknown', 'Missing, unseen or insufficient evidence. Never interpret Unknown as zero.'],
        ['Observed matches', 'Number of supplied ledgers containing records for this team, not official matches played.'],
        ['Rate', 'Total eligible visible pieces / eligible active seconds. Not maximum capability; excludes camera gaps.'],
        ['Points', 'Reviewed video-observed lower bound; unknown point values are omitted. Not official per-team totals.'],
        ['Observation eligibility', 'Direct + reviewed + not replay + not superseded. Corrections/replays remain in raw rows.'],
        ['Non-rate evidence', 'Visible motors, retention, capacity lower bounds and struggles are raw observations, not guessed totals.'],
        ['Editing', 'Read-only analytical snapshot; edit source JSON and regenerate to recompute. No macros or formulas.'],
        ['Verification', 'An export proves file generation only, not video analysis accuracy or complete coverage.']])
    # Exclusive creation also protects against a file appearing between the initial check and save.
    with output.open('xb') as handle:
        book.save(handle)
    return output

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('event', type=Path)
    parser.add_argument('ledgers', nargs='*', type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    read = lambda p: json.loads(p.read_text(encoding='utf-8-sig'))
    print(export_event(read(args.event), [read(p) for p in args.ledgers], args.output))

if __name__ == '__main__':
    try: main()
    except (ValueError, OSError, TypeError, KeyError) as error: raise SystemExit(str(error))
